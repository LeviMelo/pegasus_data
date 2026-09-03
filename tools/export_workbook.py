"""Export a built artifact to a presentable Excel workbook.

Why this exists rather than "download the CSV": a spreadsheet handed to a
team gets read without the interface's guardrails, so the workbook has to
carry them itself.

Two rules shape everything here.

**Every cut is its own aggregation.** A sheet broken down by sex is computed
by asking the aggregate layer for the sex breakdown -- never by pivoting or
re-summing another sheet. This is not fastidiousness: `los` (permanencia
media) is a MEAN, stored as the pair (los_sum, los_n) and finalised as their
quotient. Summing a column of means, or averaging them, produces a number
that looks reasonable and is wrong. Totals rows are likewise separate calls
at the coarser level, which is the only way a mean's total can be right.

**Absence is not zero.** A municipality with no cell was not observed; it did
not report zero admissions. Blank cells are left blank, and the Leia-me sheet
says so, because a reader who sees 0 will average it.

Usage::

    python tools/export_workbook.py                       # SIH, default
    python tools/export_workbook.py --artifact sim_do_municipality_month
    python tools/export_workbook.py --out C:\\path\\book.xlsx
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import urllib.request
from pathlib import Path

import pyarrow as pa
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from pegasus_data._aggregate import aggregate  # noqa: E402

API = "http://127.0.0.1:8000"

# ---------------------------------------------------------------- palette
INK = "1F2933"
ACCENT = "2F3E9E"
HEAD_FILL = PatternFill("solid", fgColor="EEF1FA")
TOTAL_FILL = PatternFill("solid", fgColor="F5F5F5")
HAIR = Side(style="thin", color="D8DEE9")


def api_json(path: str) -> dict:
    with urllib.request.urlopen(f"{API}{path}", timeout=300) as r:
        return json.load(r)


def number_format(measure: dict) -> str:
    """Excel format from the measure's declared unit and decimals."""
    if measure["unit"] == "brl":
        return 'R$ #,##0.00'
    decimals = int(measure.get("decimals") or 0)
    return "#,##0" if decimals == 0 else "#,##0." + "0" * decimals


def rows_of(table: pa.Table) -> list[dict]:
    return table.to_pylist()


def cut(
    artifact: str,
    by: list[str],
    measures: list[dict],
    time_grain: str = "month",
) -> list[dict]:
    """One cut, honouring each measure's rule for collapsing time.

    Flows (no `time_reducer`) are handed straight to the aggregate layer,
    which finalises them correctly -- a mean like `los` included.

    Stocks are the case that made the first version of this script crash on
    CNES. Beds are a quantity held AT a moment; summing them across twelve
    months yields bed-months, so the backend refuses the roll-up outright.
    The honest answer is to keep time as an axis, let each period finalise on
    its own, and then reduce those period values by the rule the artifact
    declares (mean, last or max) -- the same thing the frontend does.

    Flows and stocks therefore travel as two separate aggregations joined on
    the requested keys, because no single request is legitimate for both.
    """
    flows = [m for m in measures if not m.get("time_reducer")]
    stocks = [m for m in measures if m.get("time_reducer")]
    time_on_axis = any(b in ("month", "year", time_grain) for b in by)

    merged: dict[tuple, dict] = {}

    def absorb(rows: list[dict], value_keys: list[str]) -> None:
        for row in rows:
            key = tuple(row.get(b) for b in by)
            slot = merged.setdefault(key, {b: row.get(b) for b in by})
            for k in value_keys:
                slot[k] = row.get(k)

    if flows:
        ids = [m["id"] for m in flows]
        absorb(rows_of(aggregate(artifact, by=by or None, measures=ids)), ids)

    if stocks:
        ids = [m["id"] for m in stocks]
        if time_on_axis:
            absorb(rows_of(aggregate(artifact, by=by, measures=ids)), ids)
        else:
            per_period = rows_of(
                aggregate(artifact, by=[*by, time_grain], measures=ids)
            )
            # `last` is only meaningful in period order, which the aggregate
            # layer does not promise.
            per_period.sort(key=lambda r: str(r.get(time_grain) or ""))
            buckets: dict[tuple, dict[str, list[float]]] = {}
            for row in per_period:
                key = tuple(row.get(b) for b in by)
                slot = buckets.setdefault(key, {i: [] for i in ids})
                for i in ids:
                    v = row.get(i)
                    if v is not None:
                        slot[i].append(v)
            reduced = []
            for key, values in buckets.items():
                out = dict(zip(by, key))
                for m in stocks:
                    series = values[m["id"]]
                    if not series:
                        out[m["id"]] = None
                        continue
                    rule = m["time_reducer"]
                    if rule == "mean":
                        out[m["id"]] = sum(series) / len(series)
                    elif rule == "last":
                        out[m["id"]] = series[-1]
                    elif rule == "max":
                        out[m["id"]] = max(series)
                    else:
                        # An undeclared rule is a contract gap, not a licence
                        # to sum. Blank is the honest cell.
                        out[m["id"]] = None
                reduced.append(out)
            absorb(reduced, ids)

    return list(merged.values())


def headers_for(measures: list[dict], time_on_axis: bool) -> list[str]:
    """Measure labels, qualified when a stock has had its time collapsed.

    On a sheet with time as an axis a stock's cell IS that period's value and
    needs no qualifier. On a sheet without one, the cell is a reduction over
    the whole window, and a bare "Leitos" over 78.594 reads as a total when
    it is a monthly mean.
    """
    words = {"mean": "média", "last": "último", "max": "máximo"}
    out = []
    for m in measures:
        rule = m.get("time_reducer")
        if rule and not time_on_axis:
            out.append(f"{m['label']} ({words.get(rule, rule)} por período)")
        else:
            out.append(m["label"])
    return out


def sheet_title(name: str) -> str:
    """Excel forbids : \\ / ? * [ ] in sheet names, and caps them at 31 chars.

    "Raça/cor" is a real dimension label, so this is a certainty rather than
    a defensive flourish.
    """
    for bad in ':\\/?*[]':
        name = name.replace(bad, "-")
    return name[:31].strip() or "Planilha"


def write_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    formats: dict[int, str],
    note: str = "",
    total_row: list | None = None,
) -> None:
    """One sheet: a note, a header band, the data, an optional total."""
    ws = wb.create_sheet(sheet_title(title))
    r = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(size=9, italic=True, color="6B7280")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        r = 3

    header_row = r
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color=INK, size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=HAIR)
    r += 1

    for row in rows:
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in formats:
                cell.number_format = formats[c]
        r += 1

    if total_row is not None:
        for c, v in enumerate(total_row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
            cell.border = Border(top=HAIR)
            if c in formats:
                cell.number_format = formats[c]

    # Widths from content, capped so a long municipality name cannot push the
    # measures off screen.
    for c, h in enumerate(headers, start=1):
        longest = len(str(h))
        for row in rows[:400]:
            if c - 1 < len(row) and row[c - 1] is not None:
                longest = max(longest, len(str(row[c - 1])))
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 3, 11), 42)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
        )


def build(artifact: str, out: Path) -> Path:
    caps = api_json(f"/api/v1/datasets/{artifact}/capabilities")
    measures = caps["measures"]
    mids = [m["id"] for m in measures]

    # Territory names, indexed by the 6-digit code the artifact stores.
    membership = api_json("/api/v1/geo/membership")["membership"]
    by6: dict[str, dict] = {}
    uf_name: dict[str, str] = {}
    for entry in membership.values():
        by6[entry["code6"]] = entry
        uf_name[entry["uf"]] = entry["uf_name"]

    # Dimension labels, straight from the artifact's own codelists.
    codelists: dict[str, dict[str, str]] = {}
    for dim in caps["dimensions"]:
        payload = api_json(f"/api/v1/artifacts/{artifact}?by={dim['id']}&measures={mids[0]}")
        codelists[dim["id"]] = {
            c["code"]: c["label"] for c in (payload.get("codelists") or {}).get(dim["id"], [])
        }

    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------ Leia-me
    ws = wb.create_sheet("Leia-me")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96
    row = 1

    def kv(k: str, v: str, bold: bool = False) -> None:
        nonlocal row
        a = ws.cell(row=row, column=1, value=k)
        a.font = Font(bold=True, size=10, color=INK)
        b = ws.cell(row=row, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            b.font = Font(bold=True)
        row += 1

    def gap(n: int = 1) -> None:
        nonlocal row
        row += n

    def head(text: str) -> None:
        nonlocal row
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(bold=True, size=12, color=ACCENT)
        row += 1

    title = ws.cell(row=row, column=1, value=caps["label"])
    title.font = Font(bold=True, size=16, color=INK)
    row += 1
    ws.cell(row=row, column=1, value=caps.get("description") or "").font = Font(
        size=10, color="6B7280"
    )
    row += 2

    head("Procedência")
    kv("Fonte", f"DATASUS {caps['dataset']} ({caps['system']} · {caps['series']})")
    kv("Artefato", caps["id"])
    kv("Impressão digital", caps["fingerprint"])
    kv("Vintage (build)", caps["vintage"])
    period = caps["period"]
    kv("Período", f"{period['from']} a {period['to']} (grão: {period['grain']})")
    kv("Unidade observada", caps["observation_unit"])
    kv("Exportado em", dt.datetime.now().astimezone().isoformat(timespec="seconds"))
    gap()

    head("Medidas")
    for m in measures:
        kind = {
            "count": "contagem",
            "sum": "soma",
            "mean": "média",
            "ratio": "razão",
        }.get(m["kind"], m["kind"])
        extra = ""
        if m["kind"] == "mean":
            extra = (
                "  ⚠ É uma MÉDIA: não some nem tire média desta coluna. "
                "Cada total foi recalculado a partir dos dados de origem."
            )
        if m.get("time_reducer"):
            words = {"mean": "média", "last": "último valor", "max": "máximo"}
            extra += (
                f"  ⚠ É um ESTOQUE: existe em cada instante, não se acumula. "
                f"Somar ao longo dos meses produziria '{m['unit']}-mês'. "
                f"Onde o tempo foi colapsado, a coluna traz o "
                f"{words.get(m['time_reducer'], m['time_reducer'])} dos valores mensais."
            )
        kv(m["label"], f"{kind} · unidade: {m['unit']} · fórmula: {m['formula']}{extra}")
    gap()

    head("Como ler")
    kv(
        "Cada aba é um recorte próprio",
        "Nenhuma aba foi derivada de outra. Cada uma foi recalculada a partir "
        "do agregado, porque somar ou tirar média de uma coluna que já é média "
        "(permanência) produz um número plausível e errado.",
    )
    kv(
        "Célula vazia ≠ zero",
        "Um município sem célula não foi observado naquele recorte — ele não "
        "reportou zero. Vazio foi mantido vazio de propósito; preencher com 0 "
        "faria qualquer média subsequente mentir.",
    )
    kv(
        "Totais",
        "As linhas de Total vêm de uma agregação separada no nível mais amplo, "
        "não da soma da coluna acima.",
    )
    cov = caps["spatial"]["coverage"]
    kv(
        "Cobertura",
        f"{cov['kind']} · {cov.get('municipalities', '?')} municípios no escopo do build."
        + (f" Observação: {cov['note']}" if cov.get("note") else ""),
    )
    gap()
    kv(
        "Gerado por",
        "PegaSUS — pegasus_data (agregação) + tools/export_workbook.py. "
        "Os valores aqui são os mesmos servidos ao pegasus_view.",
    )

    # ------------------------------------------------- Resumo (Brasil/ano)
    national = cut(artifact, [], measures)
    if national:
        n = national[0]
        write_sheet(
            wb,
            "Resumo Brasil",
            ["Medida", "Valor", "Unidade"],
            [
                [h, n.get(m["id"]), m["unit"]]
                for h, m in zip(headers_for(measures, False), measures)
            ],
            {2: "#,##0.00"},
            note=f"Brasil · {period['from']} a {period['to']}. Uma agregação nacional única.",
        )
        # Per-measure formats differ, so set them cell by cell.
        sheet = wb["Resumo Brasil"]
        for i, m in enumerate(measures):
            sheet.cell(row=4 + i, column=2).number_format = number_format(m)

    # ------------------------------------------------------ tempo (mês)
    monthly = cut(artifact, ["month"], measures)
    monthly.sort(key=lambda r: r.get("month") or "")
    write_sheet(
        wb,
        "Por mês",
        ["Mês"] + headers_for(measures, True),
        [[r.get("month")] + [r.get(i) for i in mids] for r in monthly],
        {i + 2: number_format(m) for i, m in enumerate(measures)},
        note="Brasil, por mês de competência.",
        total_row=["Total"] + [national[0].get(i) for i in mids] if national else None,
    )

    # ------------------------------------------------------------- por UF
    per_uf = cut(artifact, ["uf"], measures)
    per_uf.sort(key=lambda r: -(r.get(mids[0]) or 0))
    write_sheet(
        wb,
        "Por UF",
        ["UF", "Estado"] + headers_for(measures, False),
        [
            [r.get("uf"), uf_name.get(r.get("uf"), r.get("uf"))] + [r.get(i) for i in mids]
            for r in per_uf
        ],
        {i + 3: number_format(m) for i, m in enumerate(measures)},
        note="Ordenado pela primeira medida. Cada linha é uma agregação da UF inteira.",
        total_row=["", "Total"] + [national[0].get(i) for i in mids] if national else None,
    )

    # -------------------------------------------------------- UF x mês
    uf_month = cut(artifact, ["uf", "month"], measures)
    uf_month.sort(key=lambda r: (r.get("uf") or "", r.get("month") or ""))
    write_sheet(
        wb,
        "Por UF e mês",
        ["UF", "Estado", "Mês"] + headers_for(measures, True),
        [
            [r.get("uf"), uf_name.get(r.get("uf"), r.get("uf")), r.get("month")]
            + [r.get(i) for i in mids]
            for r in uf_month
        ],
        {i + 4: number_format(m) for i, m in enumerate(measures)},
        note="Série mensal por estado.",
    )

    # ------------------------------------------------------ municípios
    per_mun = cut(artifact, ["municipality"], measures)
    per_mun.sort(key=lambda r: -(r.get(mids[0]) or 0))
    mun_rows = []
    for r in per_mun:
        code = r.get("municipality") or ""
        e = by6.get(code, {})
        mun_rows.append(
            [code, e.get("name", ""), e.get("uf", ""), e.get("macroregion", "")]
            + [r.get(i) for i in mids]
        )
    write_sheet(
        wb,
        "Por município",
        ["Código", "Município", "UF", "Região"] + headers_for(measures, False),
        mun_rows,
        {i + 5: number_format(m) for i, m in enumerate(measures)},
        note=(
            f"{len(mun_rows):,} municípios com pelo menos uma célula no período. "
            "Um município ausente desta lista não foi observado — não é zero."
        ).replace(",", "."),
        total_row=["", "", "", "Total"] + [national[0].get(i) for i in mids] if national else None,
    )

    # ------------------------------------------------- uma aba por dimensão
    for dim in caps["dimensions"]:
        did, dlabel = dim["id"], dim["label"]
        table = cut(artifact, [did], measures)
        labels = codelists.get(did, {})
        table.sort(key=lambda r: -(r.get(mids[0]) or 0))
        rows = [
            [r.get(did), labels.get(r.get(did), r.get(did))] + [r.get(i) for i in mids]
            for r in table
        ]
        note = f"Brasil, por {dlabel.lower()}."
        if dim.get("derived"):
            note += " Dimensão derivada — construída pelo pipeline, não é uma coluna bruta."
        if dim.get("note"):
            note += f" {dim['note']}"
        write_sheet(
            wb,
            f"Por {dlabel}",
            ["Código", dlabel] + headers_for(measures, False),
            rows,
            {i + 3: number_format(m) for i, m in enumerate(measures)},
            note=note,
            total_row=["", "Total"] + [national[0].get(i) for i in mids] if national else None,
        )

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--artifact", default="sih_rd_municipality_month")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    out = Path(args.out) if args.out else Path.home() / "Desktop" / f"{args.artifact}.xlsx"
    print(f"building {args.artifact} -> {out}")
    path = build(args.artifact, out)
    size = path.stat().st_size / 1_048_576
    print(f"done: {path} ({size:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
